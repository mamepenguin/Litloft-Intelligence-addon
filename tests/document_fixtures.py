"""Builders for real PDF, docx, xlsx and pptx test files."""

from __future__ import annotations

from pathlib import Path

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
EPUB = "application/epub+zip"


def make_pdf(path: Path, pages: list[str], *, user_password: str | None = None) -> Path:
    """One page per string, every line in the same font and size, so no heading is detected."""
    import fitz

    doc = fitz.open()
    for body in pages:
        page = doc.new_page()
        y = 72
        for line in body.split("\n"):
            page.insert_text((72, y), line, fontsize=11)
            y += 16
    if user_password is None:
        doc.save(str(path))
    else:
        doc.save(
            str(path),
            encryption=fitz.PDF_ENCRYPT_AES_256,
            user_pw=user_password,
            owner_pw=user_password + "-owner",
        )
    doc.close()
    return path


def make_docx(
    path: Path,
    *,
    heading: str | None = None,
    paragraphs: tuple[str, ...] = (),
    table: tuple[tuple[str, ...], ...] = (),
) -> Path:
    import docx

    document = docx.Document()
    if heading is not None:
        document.add_heading(heading, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                grid.cell(r, c).text = cell
    document.save(str(path))
    return path


def make_xlsx(path: Path, sheets: list[tuple[str, list[list[str | None]] | None]]) -> Path:
    """``None`` rows make a chartsheet charting the first worksheet; ``[]`` an empty worksheet."""
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        if rows is not None:
            sheet = workbook.create_sheet(name)
            for row in rows:
                sheet.append(row)
    for position, (name, rows) in enumerate(sheets):
        if rows is None:
            # openpyxl cannot read back a chartsheet that holds no chart.
            chartsheet = workbook.create_chartsheet(name, position)
            chart = BarChart()
            chart.add_data(Reference(workbook.worksheets[0], min_col=1, min_row=1, max_row=1))
            chartsheet.add_chart(chart)
    workbook.save(str(path))
    return path


def make_pptx(path: Path, slides: list[tuple[str | None, list[str]]]) -> Path:
    """Each slide is (title, body paragraphs). A ``None`` title uses a layout without one."""
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    title_and_content = deck.slide_layouts[1]
    blank = deck.slide_layouts[6]
    for title, body in slides:
        if title is None:
            slide = deck.slides.add_slide(blank)
            if body:
                box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
                frame = box.text_frame
                frame.text = body[0]
                for text in body[1:]:
                    frame.add_paragraph().text = text
            continue
        slide = deck.slides.add_slide(title_and_content)
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        if body:
            frame.text = body[0]
            for text in body[1:]:
                frame.add_paragraph().text = text
    deck.save(str(path))
    return path


def sparse_file(path: Path, size: int) -> Path:
    with open(path, "wb") as f:
        f.truncate(size)
    return path
