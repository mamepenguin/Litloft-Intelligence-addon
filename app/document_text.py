"""Read a document's text one unit (page, slide, sheet, section) at a time.

Units are extracted from the file on demand with the indexing extractors'
own text builders, so a unit reads as the text indexing chunked for it.
"""

from __future__ import annotations

import logging
import threading
import time
import zipfile
from dataclasses import dataclass, replace
from pathlib import Path

import app.config as config
from app.extractors import epub, office, pdf

logger = logging.getLogger(__name__)

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
EPUB = "application/epub+zip"


@dataclass(frozen=True)
class Kind:
    unit: str
    default_count: int
    max_count: int


KINDS: dict[str, Kind] = {
    PDF: Kind("page", 5, 10),
    PPTX: Kind("slide", 50, 200),
    XLSX: Kind("sheet", 5, 50),
    EPUB: Kind("section", 10, 100),
    DOCX: Kind("document", 1, 1),
}

SLOTS = 2
SLOT_WAIT_SECONDS = 10.0
_slots = threading.BoundedSemaphore(SLOTS)


class Unreadable(Exception):
    pass


class Busy(Exception):
    pass


@dataclass(frozen=True)
class Unit:
    number: int
    title: str | None
    text: str
    truncated: bool
    extractor: str


@dataclass(frozen=True)
class Extracted:
    unit_count: int
    units: tuple[Unit, ...]
    # The first requested unit the EPUB byte budget left unread, if any.
    stopped_before: int | None = None


@dataclass(frozen=True)
class Page:
    units: tuple[Unit, ...]
    next_start: int | None
    next_offset: int


def extract_units(mime: str, file_path: str, start: int, end: int) -> Extracted:
    """Units ``start`` to ``end`` (clamped to the document), holding a reader slot."""
    deadline = time.monotonic() + SLOT_WAIT_SECONDS
    if not _slots.acquire(timeout=SLOT_WAIT_SECONDS):
        raise Busy
    try:
        return _EXTRACTORS[mime](file_path, start, end, deadline)
    finally:
        _slots.release()


def page_units(extracted: Extracted, start: int, end: int, offset: int, max_chars: int) -> Page:
    """Apply ``offset`` and ``max_chars`` to the extracted range and work out the cursor."""
    last = min(end, extracted.unit_count)
    remaining = max_chars
    out: list[Unit] = []
    for index, unit in enumerate(extracted.units):
        skip = offset if index == 0 else 0
        if skip > 0 and skip >= len(unit.text):
            continue
        body = unit.text[skip:]
        if len(body) > remaining:
            out.append(replace(unit, text=body[:remaining]))
            return Page(tuple(out), unit.number, skip + remaining)
        out.append(replace(unit, text=body))
        remaining -= len(body)
        if remaining == 0:
            return _page_after(out, unit.number, extracted.unit_count)
    if extracted.stopped_before is not None:
        return Page(tuple(out), extracted.stopped_before, 0)
    return _page_after(out, last, extracted.unit_count)


def _page_after(out: list[Unit], number: int, unit_count: int) -> Page:
    if number >= unit_count:
        return Page(tuple(out), None, 0)
    return Page(tuple(out), number + 1, 0)


def _checked_path(file_path: str, max_bytes: int | None) -> Path:
    if not config.validate_file_path(file_path):
        raise Unreadable("path not allowed")
    path = Path(file_path)
    if not path.is_file():
        raise Unreadable("missing")
    if max_bytes is not None and path.stat().st_size > max_bytes:
        raise Unreadable("over the size cap")
    return path


def _pdf_units(file_path: str, start: int, end: int, deadline: float) -> Extracted:
    path = _checked_path(file_path, pdf.MAX_PDF_SIZE)
    if not pdf.PYMUPDF_LOCK.acquire(timeout=max(deadline - time.monotonic(), 0)):
        raise Busy
    try:
        return _pdf_range(path, start, end)
    finally:
        pdf.PYMUPDF_LOCK.release()


def _pdf_range(path: Path, start: int, end: int) -> Extracted:
    import fitz

    with fitz.open(str(path)) as doc:
        if doc.needs_pass:
            raise Unreadable("password-protected")
        unit_count = doc.page_count
        units = tuple(
            _pdf_page(path, doc, number)
            for number in range(start, min(end, unit_count) + 1)
        )
    return Extracted(unit_count, units)


def _pdf_page(path: Path, doc, number: int) -> Unit:
    import pymupdf4llm

    try:
        pages = pymupdf4llm.to_markdown(str(path), pages=[number - 1], page_chunks=True)
        text = pdf._page_text(list(pages)[0]).strip() if pages else ""
        extractor = pdf.EXTRACTOR_PRIMARY
    except Exception as e:
        logger.warning(
            "document-text: page %d of %s fell back to plain text (%s)",
            number, path.name, type(e).__name__,
        )
        text = (doc[number - 1].get_text("text") or "").strip()
        extractor = pdf.EXTRACTOR_FALLBACK
    return Unit(number, None, text, False, extractor)


def _docx_units(file_path: str, start: int, end: int, _deadline: float) -> Extracted:
    path = _checked_path(file_path, office.MAX_DOCX_BYTES)
    text, _heading = office.docx_text(str(path))
    units = (Unit(1, None, text, False, office.EXTRACTOR_NAME),) if start <= 1 <= end else ()
    return Extracted(1, units)


def _xlsx_units(file_path: str, start: int, end: int, _deadline: float) -> Extracted:
    import openpyxl

    path = _checked_path(file_path, office.MAX_XLSX_BYTES)
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        units: list[Unit] = []
        for number in range(start, min(end, len(names)) + 1):
            name = names[number - 1]
            ws = wb[name]
            # A chartsheet has no cells.
            rows, truncated = office.sheet_rows(ws) if hasattr(ws, "iter_rows") else ([], False)
            units.append(Unit(number, name, "\n".join(rows), truncated, office.EXTRACTOR_NAME))
    finally:
        wb.close()
    return Extracted(len(names), tuple(units))


def _pptx_units(file_path: str, start: int, end: int, _deadline: float) -> Extracted:
    from pptx import Presentation

    path = _checked_path(file_path, office.MAX_PPTX_BYTES)
    slides = list(Presentation(str(path)).slides)
    units: list[Unit] = []
    for number in range(start, min(end, len(slides)) + 1):
        title, lines = office.slide_text(slides[number - 1])
        units.append(
            Unit(number, title or None, "\n".join(lines), False, office.EXTRACTOR_NAME)
        )
    return Extracted(len(slides), tuple(units))


def _epub_units(file_path: str, start: int, end: int, _deadline: float) -> Extracted:
    path = _checked_path(file_path, None)
    with zipfile.ZipFile(path) as zf:
        entries = {epub.decode_zip_filename(info): info for info in zf.infolist()}
        budget = epub._ByteBudget(epub.BOOK_MAX_BYTES)
        book = epub._read_book(zf, entries, budget)
        if book is None:
            raise Unreadable("no container or package document")
        book = replace(book, encrypted=epub._encrypted_paths(zf, entries, budget))
        unit_count = len(book.sections)
        units: list[Unit] = []
        for number in range(start, min(end, unit_count) + 1):
            section = book.sections[number - 1]
            over = budget.spent or _declared_size(entries, book, section) > budget.left
            if over and units:
                return Extracted(unit_count, tuple(units), stopped_before=number)
            text = "" if over else epub.section_text(zf, entries, book, section, budget)
            units.append(Unit(number, None, text, False, epub.EXTRACTOR_NAME))
    return Extracted(unit_count, tuple(units))


def _declared_size(entries: dict[str, zipfile.ZipInfo], book, section) -> int:
    """Bytes reading ``section`` would take from the budget, by its zip header."""
    if section.is_nav or not section.is_html or not book.readable(section.path):
        return 0
    info = entries.get(section.path)
    return 0 if info is None else min(info.file_size, epub.SECTION_MAX_BYTES + 1)


_EXTRACTORS = {
    PDF: _pdf_units,
    DOCX: _docx_units,
    XLSX: _xlsx_units,
    PPTX: _pptx_units,
    EPUB: _epub_units,
}
